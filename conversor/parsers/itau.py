from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from ..models import AccountMeta, Transaction, ConversionResult
from ..utils import clean_text, parse_date, parse_decimal, split_account, dc_from_value


def parse(data: bytes) -> ConversionResult:
    wb=load_workbook(BytesIO(data), data_only=True, read_only=False)
    ws=wb.worksheets[0]
    agency=""; raw_account=""; header_row=None
    for r in range(1,min(ws.max_row,40)+1):
        k=clean_text(ws.cell(r,1).value).lower()
        if k.startswith("agência") or k.startswith("agencia"): agency=clean_text(ws.cell(r,2).value)
        if k.startswith("conta"): raw_account=clean_text(ws.cell(r,2).value)
        if k=="data" and "lan" in clean_text(ws.cell(r,2).value).lower(): header_row=r
    account,digit=split_account(raw_account)
    # O CIGAM da MegaThorra usa a conta Itaú sem zeros à esquerda.
    account = account.lstrip("0") or "0"
    meta=AccountMeta("Itau", agency=agency, account=account, digit=digit, source_format="xlsx")
    if not header_row: raise ValueError("Cabeçalho de lançamentos do Itaú não encontrado.")
    tx=[]
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        vals=list(row)+[None]*6
        date, lanc, razao, cpf, value, balance=vals[:6]
        if not date or not lanc: continue
        # Preserva os espaços internos do histórico exatamente como vêm do Itaú.
        desc = str(lanc).strip()
        val=parse_decimal(value); bal=parse_decimal(balance)
        # Linhas de saldo (saldo anterior e saldo total disponível do dia)
        # não possuem Valor no XLSX, mas são importadas no CIGAM com VALOR=0.
        if val is None and bal is not None:
            val=Decimal("0")
        if val is None: continue
        tx.append(Transaction(parse_date(date), desc, "", dc_from_value(val), val, bal if bal is not None else Decimal("0")))
    return ConversionResult(meta,tx)
