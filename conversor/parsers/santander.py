from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from ..models import AccountMeta, Transaction, ConversionResult
from ..utils import clean_text, parse_date, parse_decimal, split_account, dc_from_value


def parse(data: bytes) -> ConversionResult:
    wb=load_workbook(BytesIO(data), data_only=True, read_only=True)
    ws=wb.worksheets[0]
    agency=clean_text(ws.cell(1,2).value)
    raw_account=clean_text(ws.cell(1,4).value)
    account,digit=split_account(raw_account)
    meta=AccountMeta("Santander", agency=agency, account=account, digit=digit, source_format="xlsx")
    tx=[]
    header_row=None
    for r in range(1, min(ws.max_row,30)+1):
        if clean_text(ws.cell(r,1).value).lower()=="data" and "hist" in clean_text(ws.cell(r,2).value).lower():
            header_row=r; break
    if not header_row:
        raise ValueError("Cabeçalho de lançamentos do Santander não encontrado.")
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        date, hist, doc, value, balance = (list(row)+[None]*5)[:5]
        if not date or not hist:
            continue
        val=parse_decimal(value)
        bal=parse_decimal(balance)
        if val is None and "saldo" not in clean_text(hist).lower():
            continue
        if val is None: val=Decimal("0")
        tx.append(Transaction(parse_date(date), clean_text(hist), clean_text(doc) or "-", dc_from_value(val), val, bal))
    return ConversionResult(meta, tx)
